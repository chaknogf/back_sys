"""Pruebas unitarias puras del módulo Quirófano (sin DB ni HTTP).

Cubren funciones auxiliares de `reportes_service` (tiempos, promedios, KPIs,
resolución de jornada, mapeo de filas, validaciones) y el generador de Excel
`reportes_excel` (plantilla, fórmulas, validaciones, hojas dinámicas e imágenes).
"""

import io
from datetime import date, time
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException

from modules.quirofano import reportes_excel as rx
from modules.quirofano import reportes_service as rs


# ══════════════════════════════════════════════════════════
# reportes_service — helpers de tiempo
# ══════════════════════════════════════════════════════════
class TestTiempos:
    def test_t2s(self):
        assert rs._t2s(None) is None
        assert rs._t2s(time(0, 0, 0)) == 0
        assert rs._t2s(time(1, 2, 3)) == 3723

    def test_diff_seg_normal(self):
        assert rs._diff_seg(time(8, 0), time(10, 30)) == 9000

    def test_diff_seg_cruza_medianoche(self):
        assert rs._diff_seg(time(23, 0), time(1, 0)) == 7200

    def test_diff_seg_nulos(self):
        assert rs._diff_seg(None, time(10, 0)) is None
        assert rs._diff_seg(time(10, 0), None) is None

    def test_fmt_hora(self):
        assert rs._fmt_hora(None) is None
        assert rs._fmt_hora(time(7, 5)) == "07:05"

    def test_promedio_excluye_ceros_y_nulos(self):
        assert rs._promedio([0, 3600, 7200]) == 5400
        assert rs._promedio([None, 0]) is None
        assert rs._promedio([]) is None


# ══════════════════════════════════════════════════════════
# reportes_service — fila del detalle
# ══════════════════════════════════════════════════════════
def _fila_servicio() -> dict:
    interv = SimpleNamespace(
        intervencion_id=10,
        quirofano_numero_id=1,
        fecha=date(2099, 6, 1),
        paciente_id=5,
        procedimiento_principal="Apendicectomía",
        procedimiento_2=None,
        procedimiento_3=None,
        procedimiento_4=None,
        procedimiento_5=None,
        area_cuerpo_intervenida=None,
        id_area_cuerpo_intervenida=None,
        observaciones="sin novedad",
        hora_inicio_anestesia=time(7, 45),
        hora_inicio_intervencion=time(8, 0),
        hora_finaliza_intervencion=time(10, 0),
        hora_finaliza_limpieza_prepara_quirofano=time(10, 30),
        estado_cirugia_id=1,
    )
    paciente = SimpleNamespace(nombre_completo="Ana Pérez")
    medico = SimpleNamespace(nombre="Dr. López")
    estado = SimpleNamespace(nombre="Finalizada")
    formato = SimpleNamespace(nombre="Abierta")
    procedencia = SimpleNamespace(nombre="Emergencia")
    rango = SimpleNamespace(nombre="Especialista")
    quirofano = SimpleNamespace(nombre="Quirófano #1")
    especialidad = SimpleNamespace(nombre="Cirugía General")
    area = SimpleNamespace(nombre="Abdomen")
    return (interv, paciente, medico, estado, formato, procedencia, rango,
            quirofano, especialidad, area, "Apendicectomía", None, None, None, None)


class TestFilaDetalle:
    def test_mapea_campos_y_tiempos(self):
        fila = rs._fila_desde_registro(_fila_servicio(), 7)
        assert fila["no"] == 7
        assert fila["intervencion_id"] == 10
        assert fila["paciente_nombre"] == "Ana Pérez"
        assert fila["estado"] == "Finalizada"
        assert fila["procedimientos"] == "Apendicectomía"
        assert fila["hora_inicio_anestesia"] == "07:45"
        assert fila["tiempo_intervencion_seg"] == 7200   # 08:00–10:00
        assert fila["tiempo_uso_quirofano_seg"] == 9000   # 08:00–10:30
        assert fila["tiempo_limpieza_seg"] == 1800        # 10:00–10:30
        assert fila["_proc_list"] == ["Apendicectomía"]
        assert fila["area_cuerpo"] == "Abdomen"

    def test_usa_fallback_del_catalogo_area(self):
        row = list(_fila_servicio())
        row[0].area_cuerpo_intervenida = None
        fila = rs._fila_desde_registro(tuple(row), 1)
        assert fila["area_cuerpo"] == "Abdomen"


# ══════════════════════════════════════════════════════════
# reportes_service — KPIs
# ══════════════════════════════════════════════════════════
def _fila_kpi(estado, uso_seg, proc_count=1, fecha=None, interv_seg=3600):
    return {
        "estado": estado,
        "tiempo_uso_quirofano_seg": uso_seg,
        "tiempo_intervencion_seg": interv_seg,
        "tiempo_limpieza_seg": 600,
        "_proc_list": ["P"] * proc_count,
        "fecha": fecha or date(2099, 6, 1),
    }


class TestKpi:
    def test_conteos_y_porcentajes(self):
        filas = [
            _fila_kpi("Finalizada", 9000),
            _fila_kpi("Finalizada", 4500, fecha=date(2099, 6, 2)),
            _fila_kpi("Cancelada", None),
        ]
        kpi = rs._kpi(
            filas,
            quirofano_numero_id=1, numero=1, nombre="Quirófano #1",
            hora_apertura=time(7, 45), hora_cierre=time(12, 15),
            dias_abiertos=1, es_fallback=False,
        )
        assert kpi["total_intervenciones"] == 3
        assert kpi["finalizadas"] == 2
        assert kpi["canceladas"] == 1
        assert kpi["procedimientos_total"] == 3
        # (12:15 - 07:45) = 4.5 h = 16200 s de disponibilidad
        assert kpi["tiempo_disponible_seg"] == 16200
        assert kpi["porcentaje_uso"] == pytest.approx(13500 / 16200, abs=1e-4)
        assert kpi["tasa_cancelacion"] == pytest.approx(1 / 3, abs=1e-4)
        assert kpi["tasa_cancelacion_con_reprogramadas"] == pytest.approx(1 / 3, abs=1e-4)

    def test_intervalo_promedio_entre_fechas(self):
        filas = [
            _fila_kpi("Finalizada", 1000, fecha=date(2099, 6, 1)),
            _fila_kpi("Finalizada", 1000, fecha=date(2099, 6, 11)),
        ]
        kpi = rs._kpi(
            filas, quirofano_numero_id=1, numero=1, nombre="Q1",
            hora_apertura=time(7, 0), hora_cierre=time(16, 0),
            dias_abiertos=2, es_fallback=False,
        )
        # 10 días entre extremos / 2 procedimientos
        assert kpi["intervalos_dias_promedio"] == 5.0

    def test_sin_disponibilidad_no_rompe(self):
        kpi = rs._kpi(
            [_fila_kpi("Cancelada", None)],
            quirofano_numero_id=1, numero=1, nombre="Q1",
            hora_apertura=time(7, 0), hora_cierre=time(7, 0),  # disponibilidad 0
            dias_abiertos=0, es_fallback=True,
        )
        assert kpi["porcentaje_uso"] == 0.0
        assert kpi["tiempo_disponible_seg"] == 0


# ══════════════════════════════════════════════════════════
# reportes_service — resolución de jornada
# ══════════════════════════════════════════════════════════
class TestResolverJornada:
    def _filas(self):
        return [
            {"hora_inicio_anestesia": "08:00", "hora_finaliza_limpieza": "10:30",
             "fecha": date(2099, 6, 1)},
            {"hora_inicio_anestesia": "07:45", "hora_finaliza_limpieza": "12:15",
             "fecha": date(2099, 6, 2)},
        ]

    def test_prevalece_parametro_explicito(self):
        ap, ci, dias, fallback = rs._resolver_jornada(
            self._filas(), None, time(6, 0), time(18, 0), 5
        )
        assert (ap, ci, dias, fallback) == (time(6, 0), time(18, 0), 5, False)

    def test_usa_datos_del_periodo(self):
        ap, ci, dias, fallback = rs._resolver_jornada(self._filas(), None, None, None, None)
        assert ap == time(7, 45)   # MIN de anestesias
        assert ci == time(12, 15)  # MAX de limpiezas
        assert dias == 2           # fechas distintas
        assert fallback is False

    def test_fallback_a_configuracion(self):
        config = SimpleNamespace(hora_apertura=time(7, 0), hora_cierre=time(16, 0), dias_habiles=20)
        ap, ci, dias, fallback = rs._resolver_jornada([], config, None, None, None)
        assert (ap, ci, dias, fallback) == (time(7, 0), time(16, 0), 20, True)

    def test_cierre_invalido_ajusta_a_2359(self):
        ap, ci, _, _ = rs._resolver_jornada([], None, time(16, 0), time(8, 0), 0)
        assert ap == time(16, 0)
        assert ci == time(23, 59, 59)


# ══════════════════════════════════════════════════════════
# reportes_service — validaciones de generar_reporte
# ══════════════════════════════════════════════════════════
class TestValidacionesReporte:
    def test_fecha_desde_mayor(self):
        with pytest.raises(HTTPException) as exc:
            rs.generar_reporte(None, fecha_desde=date(2099, 6, 30), fecha_hasta=date(2099, 6, 1))
        assert exc.value.status_code == 422

    def test_periodo_excede_dias(self):
        with pytest.raises(HTTPException) as exc:
            rs.generar_reporte(None, fecha_desde=date(2090, 1, 1), fecha_hasta=date(2099, 1, 1))
        assert exc.value.status_code == 422


# ══════════════════════════════════════════════════════════
# reportes_excel — helpers
# ══════════════════════════════════════════════════════════
class TestExcelHelpers:
    def test_parse_hora(self):
        assert rx._parse_hora(None) is None
        assert rx._parse_hora("") is None
        assert rx._parse_hora("08:30") == time(8, 30)
        assert rx._parse_hora("malo") is None

    def test_rango_lista(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["AK7"] = "a"
        ws["AK8"] = "b"
        ws["AK9"] = "c"
        assert rx._rango_lista(ws, "AK") == "$AK$7:$AK$9"
        ws["AK7"] = None
        assert rx._rango_lista(ws, "AK") == "$AK$7:$AK$7"


# ══════════════════════════════════════════════════════════
# reportes_excel — construcción del libro
# ══════════════════════════════════════════════════════════
def _fila_excel(**over) -> dict:
    base = {
        "fecha": date(2099, 6, 1),
        "paciente_nombre": "Ana Pérez",
        "rango_medico": "Especialista",
        "especialidad": "Cirugía General",
        "medico": "Dr. López",
        "procedencia": "Emergencia",
        "formato": "Abierta",
        "procedimiento_1": "Apendicectomía",
        "procedimiento_2": None,
        "procedimiento_3": None,
        "procedimiento_4": None,
        "procedimiento_5": None,
        "area_cuerpo": "Abdomen",
        "observaciones": "sin novedad",
        "hora_inicio_anestesia": "07:45",
        "hora_inicio_intervencion": "08:00",
        "hora_finaliza_intervencion": "10:00",
        "hora_finaliza_limpieza": "10:30",
        "estado": "Finalizada",
    }
    base.update(over)
    return base


def _group(numero, n_filas=1, **kpi_over):
    kpi = {
        "quirofano_numero_id": numero,
        "numero": numero,
        "nombre": f"Quirófano #{numero}",
        "hora_apertura": "07:45",
        "hora_cierre": "12:15",
        "es_fallback": False,
        "dias_abiertos": 1,
        "tiempo_disponible_seg": 16200,
        "horas_activo_seg": 16200,
        "total_intervenciones": n_filas,
        "finalizadas": n_filas,
        "canceladas": 0,
        "reprogramadas": 0,
        "procedimientos_total": n_filas,
        "intervalos_dias_promedio": None,
        "tiempo_promedio_intervencion_seg": 7200,
        "tiempo_promedio_uso_quirofano_seg": 9000,
        "tiempo_promedio_limpieza_seg": 1800,
        "porcentaje_uso": 0.83,
        "tasa_cancelacion": 0.0,
        "tasa_cancelacion_con_reprogramadas": 0.0,
    }
    kpi.update(kpi_over)
    return {"kpi": kpi, "intervenciones": [_fila_excel() for _ in range(n_filas)]}


def _reporte(grupos, procedimientos=None, reprogramadas=0):
    return {
        "fecha_desde": date(2099, 6, 1),
        "fecha_hasta": date(2099, 6, 30),
        "generado_en": "2099-06-30T10:00:00",
        "filtros": {},
        "kpi_global": {"reprogramadas": reprogramadas},
        "por_quirofano": grupos,
        "procedimientos": procedimientos or [],
        "advertencias": [],
    }


class TestExcelLibro:
    def test_genera_hojas_q_y_resumen_al_final(self):
        data = rx.construir_excel_bytes(_reporte([_group(1), _group(2)]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Q1", "Q2", "Resumen"]

    def test_resumen_suma_totales_de_dos_hojas(self):
        data = rx.construir_excel_bytes(_reporte([_group(1), _group(2)]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb["Resumen"]
        assert ws["E11"].value == "='Q1'!Z22+'Q2'!Z22"
        assert ws["G11"].value == "='Q1'!AB22+'Q2'!AB22"

    def test_crea_hoja_dinamica_q3(self):
        data = rx.construir_excel_bytes(_reporte([_group(1), _group(2), _group(3)]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Q1", "Q2", "Q3", "Resumen"]
        # limitación conocida: la hoja copiada no arrastra el logo
        assert len(wb["Q3"]._images) == 0

    def test_excluye_grupo_sin_quirofano(self):
        sin_q = _group(None)
        sin_q["kpi"]["nombre"] = "Sin quirófano asignado"
        data = rx.construir_excel_bytes(_reporte([_group(1), sin_q]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Q1", "Resumen"]

    def test_imagenes_de_plantilla_preservadas(self):
        data = rx.construir_excel_bytes(_reporte([_group(1), _group(2)]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert len(wb["Q1"]._images) >= 1
        assert len(wb["Q2"]._images) >= 1

    def test_contenido_y_formulas_de_hoja(self):
        data = rx.construir_excel_bytes(_reporte([_group(1)]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb["Q1"]
        assert ws["B6"].value == "QUIRÓFANO #1"
        assert ws["D9"].value == time(7, 45)
        assert ws["D10"].value == time(12, 15)
        assert ws["D11"].value == 1
        assert ws["A14"].value == 1
        assert ws["B14"].value.date() == date(2099, 6, 1)
        assert ws["B14"].number_format == "dd/mm/yyyy"
        assert ws["I14"].value == "Apendicectomía"
        assert ws["P14"].value == time(7, 45)
        assert ws["T14"].value == "=R14-Q14"
        assert ws["U14"].value == "=S14-Q14"
        assert ws["W14"].value == "=S14-R14"
        assert "TEXTJOIN" in ws["AM14"].value
        assert ws["Z15"].value == "=D10-D9"
        assert ws["B614"].value == "Total"

    def test_validaciones_recreadas(self):
        data = rx.construir_excel_bytes(_reporte([_group(1)]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert len(wb["Q1"].data_validations.dataValidation) == 5

    def test_hoja_desbordada_quita_fila_total(self):
        data = rx.construir_excel_bytes(_reporte([_group(1, n_filas=601)]))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb["Q1"]
        assert ws["B614"].value != "Total"
        assert ws["B614"].value.date() == date(2099, 6, 1)
        assert "B614:T614" not in {str(r) for r in ws.merged_cells.ranges}

    def test_resumen_escribe_procedimientos_y_reprogramadas(self):
        procs = [{
            "procedimiento": "Apendicectomía",
            "cantidad": 2,
            "tiempo_promedio_seg": 3600,
            "porcentaje_uso": 0.1,
            "intervalo_dias_promedio": None,
        }]
        data = rx.construir_excel_bytes(_reporte([_group(1, n_filas=2)], procs, reprogramadas=3))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb["Resumen"]
        assert ws["I10"].value == "Reprogramadas"
        assert ws["I11"].value == 3
        assert ws["A8"].value == 2
        assert ws["B8"].value == "Apendicectomía"
        assert "SUMIF" in ws["E8"].value
        assert ws["E11"].value.startswith("=")
        assert ws["G11"].value.startswith("=")
        assert ws["E11"].value == "='Q1'!Z22"
        assert ws["G11"].value == "='Q1'!AB22"
