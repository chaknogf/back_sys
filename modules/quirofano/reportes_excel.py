"""Genera el libro .xlsx del reporte de quirófanos.

Parte de la plantilla `assets/templates/Control_Quirofanos_plantilla.xlsx`
(hojas Q1, Q2 y Resumen con estilos, combinadas, validaciones e imágenes) y
escribe los datos del período. Si hay más quirófanos que hojas en la plantilla
se copia la hoja `Q1` (la copia no arrastra el logo, limitación de openpyxl) y
siempre se colocan las hojas Q… antes de `Resumen`.
"""

from __future__ import annotations

from datetime import time
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

PLANTILLA = (
    Path(__file__).resolve().parents[2] / "assets" / "templates" / "Control_Quirofanos_plantilla.xlsx"
)
MAX_FILAS_HOJA = 600
_FILA_ENCABEZADO = 13
_FILA_INICIO = 14


def _parse_hora(valor: str | None) -> time | None:
    if not valor:
        return None
    try:
        hh, mm = valor.split(":")[:2]
        return time(int(hh), int(mm))
    except (ValueError, TypeError):
        return None


def _rango_lista(ws, columna: str, desde: int = 7) -> str:
    """Devuelve el rango continuo no vacío de una columna de catálogo auxiliar."""
    fila = desde
    ultima = desde
    while fila <= desde + 1000:
        if ws[f"{columna}{fila}"].value not in (None, ""):
            ultima = fila
            fila += 1
        else:
            break
    return f"${columna}${desde}:${columna}${ultima}"


def _obtener_hoja(wb, nombre: str):
    if nombre in wb.sheetnames:
        return wb[nombre]
    base = wb["Q1"] if "Q1" in wb.sheetnames else wb[wb.sheetnames[0]]
    ws = wb.copy_worksheet(base)
    ws.title = nombre
    return ws


def _escribir_hoja(ws, group: dict) -> int:
    """Escribe una hoja Q. Devuelve la última fila de datos."""
    kpi = group["kpi"]
    filas = group["intervenciones"]
    total_filas = max(MAX_FILAS_HOJA, len(filas))
    ultima = _FILA_ENCABEZADO + total_filas

    # Si hay desbordamiento, la fila de total se pierde: quitar su combinación
    # ANTES de escribir (la fila 614 pasa a ser datos y no admite celdas combinadas).
    if total_filas > MAX_FILAS_HOJA:
        for rng in list(ws.merged_cells.ranges):
            if str(rng) == "B614:T614":
                ws.unmerge_cells("B614:T614")

    ws["B6"] = kpi["nombre"].upper()
    ws["D9"] = _parse_hora(kpi["hora_apertura"])
    ws["D10"] = _parse_hora(kpi["hora_cierre"])
    ws["D11"] = kpi["dias_abiertos"]

    for i in range(1, total_filas + 1):
        r = _FILA_ENCABEZADO + i
        for c in range(1, 23):  # A..V
            ws.cell(r, c).value = None
        ws.cell(r, 1).value = i
        if i <= len(filas):
            f = filas[i - 1]
            ws.cell(r, 2).value = f["fecha"]
            ws.cell(r, 2).number_format = "dd/mm/yyyy"
            ws.cell(r, 3).value = f["paciente_nombre"]
            ws.cell(r, 4).value = f["rango_medico"]
            ws.cell(r, 5).value = f["especialidad"]
            ws.cell(r, 6).value = f["medico"]
            ws.cell(r, 7).value = f["procedencia"]
            ws.cell(r, 8).value = f["formato"]
            ws.cell(r, 9).value = f["procedimiento_1"]
            ws.cell(r, 10).value = f["procedimiento_2"]
            ws.cell(r, 11).value = f["procedimiento_3"]
            ws.cell(r, 12).value = f["procedimiento_4"]
            ws.cell(r, 13).value = f["procedimiento_5"]
            ws.cell(r, 14).value = f["area_cuerpo"]
            ws.cell(r, 15).value = f["observaciones"]
            ws.cell(r, 16).value = _parse_hora(f["hora_inicio_anestesia"])
            ws.cell(r, 17).value = _parse_hora(f["hora_inicio_intervencion"])
            ws.cell(r, 18).value = _parse_hora(f["hora_finaliza_intervencion"])
            ws.cell(r, 19).value = _parse_hora(f["hora_finaliza_limpieza"])
            ws.cell(r, 22).value = f["estado"]
        for c in (16, 17, 18, 19):  # P..S
            ws.cell(r, c).number_format = "h:mm"
        ws.cell(r, 20).value = f"=R{r}-Q{r}"  # T
        ws.cell(r, 20).number_format = "[h]:mm"
        ws.cell(r, 21).value = f"=S{r}-Q{r}"  # U
        ws.cell(r, 21).number_format = "[h]:mm"
        ws.cell(r, 23).value = f"=S{r}-R{r}"  # W
        ws.cell(r, 23).number_format = "[h]:mm"
        ws.cell(r, 39).value = f'=_xlfn.TEXTJOIN(" + ", TRUE, I{r}:M{r})'  # AM

    # Fila de total (solo cuando no hay desbordamiento de filas de datos).
    if total_filas == MAX_FILAS_HOJA:
        ws["B614"] = "Total"

    # Panel de KPIs (referencias ajustadas al rango real).
    ws["Z15"] = "=D10-D9"
    ws["Z18"] = "=(D10-D9)*D11"
    ws["AB15"] = f'=IFERROR(SUMIF(V14:V{ultima},"Finalizada",U14:U{ultima})/Z18,0)'
    ws["AD15"] = "=IFERROR(AB22/(Z22+AB22),0)"
    ws["Z22"] = f'=COUNTIF($V$14:$V${ultima},"Finalizada")'
    ws["AB22"] = f'=COUNTIF($V$14:$V${ultima},"Cancelada")'
    ws["Z25"] = f"=COUNTA(V14:V{ultima})"
    ws["Z29"] = f'=IFERROR(AVERAGEIF(T14:T{ultima},"<>0"),"")'
    ws["AC29"] = f'=IFERROR(AVERAGEIF(W14:W{ultima},"<>0"),"")'
    ws["Z32"] = f'=IFERROR(AVERAGEIF(U14:U{ultima},"<>0"),"")'

    # Validaciones de lista (recreadas con el rango de filas real).
    ws.data_validations.dataValidation = []
    reglas = [
        ("D", _rango_lista(ws, "AK")),
        ("I", _rango_lista(ws, "AJ")),
        ("G", _rango_lista(ws, "AL")),
        ("V", _rango_lista(ws, "AI")),
        ("H", _rango_lista(ws, "AH")),
    ]
    for columna, formula in reglas:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        if columna == "I":
            dv.add(f"I{_FILA_INICIO}:M{ultima}")
        else:
            dv.add(f"{columna}{_FILA_INICIO}:{columna}{ultima}")

    return ultima


def _escribir_resumen(wb, hojas: list[tuple[str, int]], procedimientos: list[dict], kpi_global: dict):
    ws = wb["Resumen"]
    nombres = [n for n, _ in hojas]

    ws["C3"] = f"='{nombres[0]}'!C3" if nombres else None

    # Limpiar tabla de procedimientos y columna auxiliar R.
    for r in range(8, 501):
        ws.cell(r, 1).value = None
        ws.cell(r, 2).value = None
        ws.cell(r, 3).value = None
    for r in range(4, 4205):
        ws.cell(r, 18).value = None

    def suma_por_hoja(columna: str, condicion: str, valor: str) -> str:
        return "+".join(
            f"SUMIF('{n}'!V14:V{last},\"{condicion}\",'{n}'!{valor}14:{valor}{last})"
            for n, last in hojas
        )

    def count_por_hoja(condicion: str) -> str:
        return "+".join(
            f'COUNTIF(\'{n}\'!V14:V{last},"{condicion}")' for n, last in hojas
        )

    uso = suma_por_hoja("V", "Finalizada", "U")
    z18 = "+".join(f"'{n}'!Z18" for n, _ in hojas)
    ws["E8"] = f"=IFERROR(({uso})/({z18}),0)"
    ws["G8"] = (
        f"=IFERROR(({count_por_hoja('Cancelada')})/"
        f"(({count_por_hoja('Finalizada')})+({count_por_hoja('Cancelada')})"
        f"+({count_por_hoja('Reprogramada')})),0)"
    )
    ws["E11"] = "=" + "+".join(f"'{n}'!Z22" for n, _ in hojas)
    ws["G11"] = "=" + "+".join(f"'{n}'!AB22" for n, _ in hojas)
    ws["I10"] = "Reprogramadas"
    ws["I11"] = kpi_global.get("reprogramadas", 0)
    ws["E14"] = "=E11+G11+I11"

    # Tabla por procedimiento (valores calculados en backend).
    formato_tiempo = ws["C8"].number_format
    for i, proc in enumerate(procedimientos):
        r = 8 + i
        ws.cell(r, 1).value = proc["cantidad"]
        ws.cell(r, 2).value = proc["procedimiento"]
        tiempo = proc.get("tiempo_promedio_seg")
        if tiempo is not None:
            ws.cell(r, 3).value = tiempo / 86400
            ws.cell(r, 3).number_format = formato_tiempo

    ultimo_proc = max(22, 8 + len(procedimientos) - 1)
    ws["E17"] = f'=IFERROR(AVERAGEIF(C8:C{ultimo_proc},">0"),"")'


def construir_excel_bytes(reporte: dict) -> bytes:
    wb = openpyxl.load_workbook(PLANTILLA)

    grupos = [g for g in reporte["por_quirofano"] if g["kpi"]["numero"] is not None]
    hojas: list[tuple[str, int]] = []

    for group in grupos:
        numero = group["kpi"]["numero"]
        nombre = f"Q{numero}"
        ws = _obtener_hoja(wb, nombre)
        ultima = _escribir_hoja(ws, group)
        hojas.append((nombre, ultima))

    # Eliminar hojas Q sobrantes de la plantilla.
    deseadas = {n for n, _ in hojas}
    for existente in list(wb.sheetnames):
        if existente.startswith("Q") and existente not in deseadas:
            del wb[existente]

    _escribir_resumen(wb, hojas, reporte["procedimientos"], reporte["kpi_global"])

    # Orden final: Q… y Resumen al final.
    wb._sheets = [wb[n] for n, _ in hojas] + [wb["Resumen"]]

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
